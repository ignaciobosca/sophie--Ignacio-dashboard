#!/usr/bin/env python3
"""Sophie Society MBR deck builder.

Reads report_data.json and renders a Sophie-branded 16:9 PPTX from its `slides`
array. Charts are referenced by filename from <json_dir>/charts/ (run
build_charts.py first). Output .pptx is written next to the json.

Usage:  python build_deck.py <path/to/report_data.json>

Rich text uses "runs": a list of [text, style] where style keys are
{size, bold, italic, color}. color is a hex ("#13835B") or a token:
DARK GREEN GOLD RED G900 G700 G300 WHITE. See references/slides.md for the
full slide-type catalog and JSON shapes.
"""
import json, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

GREEN=RGBColor(0x13,0x83,0x5B); DARK=RGBColor(0x0D,0x5C,0x3F); LIGHT=RGBColor(0xE8,0xF5,0xF0)
GOLD=RGBColor(0xF5,0xA6,0x23); G900=RGBColor(0x1A,0x1A,0x1A); G700=RGBColor(0x4A,0x4A,0x4A)
G300=RGBColor(0xD4,0xD4,0xD4); G100=RGBColor(0xF5,0xF5,0xF5); WHITE=RGBColor(0xFF,0xFF,0xFF)
REDX=RGBColor(0xC0,0x39,0x2B); REDL=RGBColor(0xFB,0xEC,0xEA); MINT=RGBColor(0x8F,0xD3,0xB8)
CREAM=RGBColor(0xC8,0xDE,0xD5)
FONT="Calibri"
SW, SH = Inches(13.333), Inches(7.5)
_TOK={"DARK":DARK,"GREEN":GREEN,"GOLD":GOLD,"RED":REDX,"G900":G900,"G700":G700,
      "G300":G300,"WHITE":WHITE,"LIGHT":LIGHT,"MINT":MINT,"CREAM":CREAM}

def col(v, default=G900):
    if v is None: return default
    if isinstance(v, RGBColor): return v
    if isinstance(v, str):
        if v in _TOK: return _TOK[v]
        h = v.lstrip("#")
        if len(h)==6:
            return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))
    return default

ACCENTS={"GOLD":GOLD,"GREEN":GREEN,"RED":REDX,"DARK":DARK}
CARDBG={"GOLD":WHITE,"GREEN":WHITE,"RED":REDL,"DARK":WHITE}

class Deck:
    def __init__(self, brand):
        self.p = Presentation(); self.p.slide_width=SW; self.p.slide_height=SH
        self.blank = self.p.slide_layouts[6]; self.brand=brand; self.n=0
    def slide(self): return self.p.slides.add_slide(self.blank)
    def rect(self,s,x,y,w,h,c,line=None):
        sp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h)
        sp.fill.solid(); sp.fill.fore_color.rgb=c
        if line is None: sp.line.fill.background()
        else: sp.line.color.rgb=line; sp.line.width=Pt(0.75)
        sp.shadow.inherit=False; return sp
    def txt(self,s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,sp_after=6,line_sp=1.06):
        tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame
        tf.word_wrap=True; tf.vertical_anchor=anchor
        tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
        if isinstance(runs,str): runs=[[[runs,{}]]]
        for i,para in enumerate(runs):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.alignment=align; p.space_after=Pt(sp_after); p.line_spacing=line_sp
            if isinstance(para,str): para=[[para,{}]]
            for run in para:
                if isinstance(run,str): text, st = run, {}
                else: text, st = (run[0], run[1] if len(run)>1 else {})
                r=p.add_run(); r.text=text; f=r.font
                f.name=st.get("font",FONT); f.size=Pt(st.get("size",14))
                f.bold=st.get("bold",False); f.italic=st.get("italic",False)
                f.color.rgb=col(st.get("color"))
        return tb
    def header(self,s,title,kicker=None):
        # No accent rules/stripes under the title — flat, table-first look (see references/slides.md design notes).
        if kicker:
            self.txt(s,Inches(0.55),Inches(0.32),Inches(11),Inches(0.32),[[[kicker,{"size":11,"bold":True,"color":GREEN}]]])
            ty=0.60
        else:
            ty=0.42
        self.txt(s,Inches(0.55),Inches(ty),Inches(12.2),Inches(0.6),[[[title,{"size":26,"bold":True,"color":GREEN}]]])
    def footer(self,s):
        self.n+=1
        self.txt(s,Inches(0.55),Inches(7.06),Inches(9),Inches(0.3),
                 [[[f"Sophie Society  ·  {self.brand}  ·  Confidential",{"size":9,"color":G300}]]])
        self.txt(s,Inches(11.8),Inches(7.06),Inches(1.0),Inches(0.3),
                 [[[str(self.n),{"size":9,"color":G300}]]],align=PP_ALIGN.RIGHT)
    def bullets(self,s,items,x,y,w,h,size=14,gap=8):
        tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True
        tf.margin_left=0; tf.margin_right=0; first=True
        for it in items:
            head=it.get("head") if isinstance(it,dict) else None
            body=it.get("body") if isinstance(it,dict) else it
            p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
            p.space_after=Pt(gap); p.line_spacing=1.05
            r=p.add_run(); r.text="▸  "; r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=GREEN; r.font.bold=True
            if head:
                r2=p.add_run(); r2.text=head+"  "; r2.font.name=FONT; r2.font.size=Pt(size); r2.font.bold=True; r2.font.color.rgb=DARK
            r3=p.add_run(); r3.text=body; r3.font.name=FONT; r3.font.size=Pt(size); r3.font.color.rgb=G700

# ---- slide renderers ----
def s_cover(dk, sd):
    s=dk.slide(); dk.rect(s,0,0,SW,SH,DARK); dk.rect(s,0,0,Inches(0.35),SH,GREEN)
    dk.rect(s,Inches(0.9),Inches(2.05),Inches(1.5),Pt(4),GOLD)
    dk.txt(s,Inches(0.9),Inches(1.25),Inches(11),Inches(0.5),[[["SOPHIE SOCIETY",{"size":15,"bold":True,"color":MINT}]]])
    dk.txt(s,Inches(0.9),Inches(2.35),Inches(11.7),Inches(1.6),
           [[[sd["brand"],{"size":46,"bold":True,"color":WHITE}]],
            [[sd.get("subtitle","Monthly Performance Report"),{"size":26,"color":CREAM}]]],sp_after=6)
    dk.txt(s,Inches(0.9),Inches(4.55),Inches(11),Inches(0.5),[[[sd.get("period_label","").upper(),{"size":20,"bold":True,"color":GOLD}]]])
    lines=[]
    if sd.get("period"): lines.append([[sd["period"],{"size":13,"color":CREAM}]])
    if sd.get("meta"):   lines.append([[sd["meta"],{"size":12,"italic":True,"color":RGBColor(0x9C,0xC4,0xB6)}]])
    if lines: dk.txt(s,Inches(0.9),Inches(5.35),Inches(11.6),Inches(1.2),lines,sp_after=5)

def s_exec(dk, sd):
    s=dk.slide(); dk.header(s, sd.get("title","Executive Summary"))
    dk.rect(s,Inches(0.55),Inches(1.4),Inches(12.25),Inches(1.35),LIGHT)
    dk.txt(s,Inches(0.85),Inches(1.52),Inches(11.7),Inches(1.1),[sd["highlight"]],anchor=MSO_ANCHOR.MIDDLE)
    tiles=sd.get("tiles",[])[:4]; tw=Inches(2.94); gap=Inches(0.17); tx=Inches(0.55)
    for i,t in enumerate(tiles):
        x=tx+i*(tw+gap); dk.rect(s,x,Inches(3.05),tw,Inches(1.6),G100)
        dk.txt(s,x,Inches(3.28),tw,Inches(0.4),[[[t["label"].upper(),{"size":11,"bold":True,"color":G700}]]],align=PP_ALIGN.CENTER)
        dk.txt(s,x,Inches(3.68),tw,Inches(0.7),[[[str(t["val"]),{"size":30,"bold":True,"color":DARK}]]],align=PP_ALIGN.CENTER)
        dk.txt(s,x,Inches(4.34),tw,Inches(0.3),[[[t.get("sub",""),{"size":11,"color":GREEN,"bold":True}]]],align=PP_ALIGN.CENTER)
    if sd.get("working"):
        dk.txt(s,Inches(0.55),Inches(5.0),Inches(12),Inches(0.4),[[[sd.get("working_title","What's working"),{"size":13,"bold":True,"color":DARK}]]])
        dk.bullets(s,[{"body":b} for b in sd["working"]],Inches(0.6),Inches(5.42),Inches(12.2),Inches(1.4),size=13,gap=6)
    dk.footer(s)

def s_kpi(dk, sd):
    s=dk.slide(); dk.header(s, sd.get("title","KPI Scorecard"))
    rows=[sd["header"]]+sd["rows"]; dirs=sd.get("delta_dir",[])
    nrow=len(rows)
    gt=s.shapes.add_table(nrow,5,Inches(0.55),Inches(1.4),Inches(12.25),Inches(5.15)).table
    gt.columns[0].width=Inches(3.55)
    for c in range(1,5): gt.columns[c].width=Inches(2.175)
    for r in range(nrow):
        gt.rows[r].height=Inches(0.5)
        for c in range(5):
            cell=gt.cell(r,c); cell.margin_left=Inches(0.12); cell.margin_right=Inches(0.08)
            cell.margin_top=Inches(0.02); cell.margin_bottom=Inches(0.02); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            p=cell.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.CENTER
            run=p.add_run(); run.text=str(rows[r][c]); run.font.name=FONT; run.font.size=Pt(13 if r==0 else 12.5)
            if r==0:
                cell.fill.solid(); cell.fill.fore_color.rgb=DARK; run.font.bold=True; run.font.color.rgb=WHITE
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb=WHITE if r%2 else G100; run.font.color.rgb=G900
                if c==0: run.font.bold=True
                if c==2: run.font.bold=True; run.font.color.rgb=DARK
                if c==3:
                    d=dirs[r-1] if r-1<len(dirs) else "neutral"; run.font.bold=True
                    run.font.color.rgb=GREEN if d=="up" else (REDX if d=="down" else G700)
    if sd.get("footnote"):
        dk.txt(s,Inches(0.55),Inches(6.7),Inches(12.2),Inches(0.4),[[[sd["footnote"],{"size":10,"italic":True,"color":G700}]]])
    dk.footer(s)

def s_chart(dk, sd, cdir):
    s=dk.slide(); dk.header(s, sd["title"])
    img=os.path.join(cdir, sd["image"])
    if os.path.exists(img):
        s.shapes.add_picture(img, Inches(0.7), Inches(1.45), width=Inches(9.5))
    if sd.get("takeaway"):
        dk.rect(s,Inches(10.4),Inches(1.7),Inches(2.5),Inches(4.6),LIGHT)
        dk.txt(s,Inches(10.62),Inches(1.95),Inches(2.1),Inches(0.4),[[["TAKEAWAY",{"size":11,"bold":True,"color":GREEN}]]])
        dk.txt(s,Inches(10.62),Inches(2.4),Inches(2.12),Inches(3.7),[[[sd["takeaway"],{"size":12.5,"color":G900}]]],line_sp=1.12)
    dk.footer(s)

def s_narr(dk, sd):
    s=dk.slide(); dk.header(s, sd["title"])
    paras=[]
    for p in sd["paragraphs"]:
        run=[]
        if p.get("head"): run.append([p["head"]+"  ",{"size":12.5,"bold":True,"color":DARK}])
        run.append([p["body"],{"size":12.5,"color":G700}])
        paras.append(run)
    dk.txt(s,Inches(0.55),Inches(1.4),Inches(12.25),Inches(5.5),paras,sp_after=9,line_sp=1.06)
    dk.footer(s)

def s_bullets(dk, sd):
    s=dk.slide(); dk.header(s, sd["title"])
    y=1.35
    if sd.get("subtitle"):
        dk.txt(s,Inches(0.55),Inches(y),Inches(12),Inches(0.4),[[[sd["subtitle"],{"size":12,"italic":True,"color":G700}]]]); y=1.9
    else: y=1.55
    hb = 4.4 if (sd.get("footnote") or sd.get("banner")) else 5.2
    dk.bullets(s,sd["items"],Inches(0.6),Inches(y),Inches(12.2),Inches(hb),
               size=sd.get("size",13.5),gap=sd.get("gap",12))
    if sd.get("banner"):
        dk.rect(s,Inches(0.55),Inches(6.0),Inches(12.25),Inches(0.85),G100)
        dk.txt(s,Inches(0.8),Inches(6.1),Inches(11.7),Inches(0.65),
               [[[sd["banner"].get("label","Bottom line:  "),{"size":13,"bold":True,"color":DARK}],
                 [sd["banner"]["text"],{"size":13,"color":G700}]]],anchor=MSO_ANCHOR.MIDDLE)
    elif sd.get("footnote"):
        dk.txt(s,Inches(0.55),Inches(6.5),Inches(12.2),Inches(0.4),[[[sd["footnote"],{"size":10,"italic":True,"color":G700}]]])
    dk.footer(s)

def _card(dk,s,x,w,c):
    acc=ACCENTS.get(c.get("accent","GREEN"),GREEN); bg=CARDBG.get(c.get("accent","GREEN"),WHITE)
    dk.rect(s,x,Inches(1.85),w,Inches(4.5),bg,line=G300 if bg==WHITE else None)
    dk.txt(s,x+Inches(0.2),Inches(2.0),w-Inches(0.4),Inches(0.4),[[[c["heading"],{"size":13,"bold":True,"color":DARK if c.get("accent")!="RED" else REDX}]]])
    yy=2.5
    if c.get("big"):
        dk.txt(s,x+Inches(0.2),Inches(2.25),w-Inches(0.4),Inches(1.0),[[[str(c["big"]),{"size":46,"bold":True,"color":acc}]]],align=PP_ALIGN.CENTER)
        dk.txt(s,x+Inches(0.2),Inches(3.35),w-Inches(0.4),Inches(0.4),[[[c.get("big_sub",""),{"size":12,"color":G700}]]],align=PP_ALIGN.CENTER)
        yy=4.0
    if c.get("bullets"):
        dk.bullets(s,c["bullets"],x+Inches(0.2),Inches(yy),w-Inches(0.4),Inches(6.2-yy),size=12.5,gap=9)
    elif c.get("body"):
        dk.txt(s,x+Inches(0.2),Inches(yy),w-Inches(0.4),Inches(6.2-yy),[[[c["body"],{"size":13,"color":G700}]]],line_sp=1.15)

def s_twocard(dk, sd):
    s=dk.slide(); dk.header(s, sd["title"])
    if sd.get("subtitle"):
        dk.txt(s,Inches(0.55),Inches(1.32),Inches(12),Inches(0.4),[sd["subtitle"]] if isinstance(sd["subtitle"],list) else [[[sd["subtitle"],{"size":13,"color":G700}]]])
    _card(dk,s,Inches(0.55),Inches(6.0),sd["left"])
    _card(dk,s,Inches(6.8),Inches(6.0),sd["right"])
    dk.footer(s)

def s_pl(dk, sd):
    s=dk.slide(); dk.header(s, sd["title"])
    y0=1.55; n=len(sd["rows"])
    reserved = (0.5 if sd.get("note") else 0) + (0.5 if sd.get("footnote") else 0)
    avail = 5.35 - reserved
    row_h=min(0.585, avail/max(n,1))
    y=Inches(y0)
    for row in sd["rows"]:
        bold=row.get("bold",False)
        dk.rect(s,Inches(0.55),y,Inches(6.6),Inches(row_h),LIGHT if bold else WHITE)
        dk.txt(s,Inches(0.72),y,Inches(4.6),Inches(row_h),[[[row["label"],{"size":12.5 if row_h>0.4 else 11,"bold":bold,"color":G900}]]],anchor=MSO_ANCHOR.MIDDLE)
        dk.txt(s,Inches(5.2),y,Inches(1.8),Inches(row_h),[[[row["val"],{"size":12.5 if row_h>0.4 else 11,"bold":bold,"color":col(row.get("color"))}]]],align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE)
        y=y+Inches(row_h)
    y_text = y+Inches(0.1)
    if sd.get("note"):
        dk.txt(s,Inches(0.72),y_text,Inches(6.4),Inches(0.5),[[[sd["note"],{"size":10,"italic":True,"color":G700}]]],line_sp=1.0)
        y_text = y_text+Inches(0.45)
    pan=sd.get("panel")
    if pan:
        acc=ACCENTS.get(pan.get("accent","GREEN"),GREEN)
        dk.rect(s,Inches(7.55),Inches(1.55),Inches(5.25),Inches(4.65),LIGHT)
        dk.txt(s,Inches(7.8),Inches(1.75),Inches(4.8),Inches(0.5),[[[pan["title"],{"size":13,"bold":True,"color":acc if pan.get("accent")=="RED" else DARK}]]])
        dk.txt(s,Inches(7.8),Inches(2.4),Inches(4.8),Inches(3.6),pan["runs"],sp_after=8,line_sp=1.15)
    if sd.get("footnote"):
        dk.txt(s,Inches(0.55),y_text,Inches(6.4),Inches(0.5),[[[sd["footnote"],{"size":9.5,"italic":True,"color":G700}]]],line_sp=1.0)
    dk.footer(s)

def _cell_text_color(cell):
    if isinstance(cell,(list,tuple)) and len(cell)==2 and not isinstance(cell[1],dict):
        return str(cell[0]), col(cell[1])
    return str(cell), G900

def s_data_table(dk, sd):
    """Table-first finding slide: short intro (optional) -> table (green header, zebra rows,
    cells may be 'text' or ['text','COLOR_TOKEN'] for flagged values) -> optional side chart_image
    -> caption paragraph -> optional bottom banner. Mirrors the AdLabs-audit look: data first,
    prose kept to 1-3 short lines above/below, no decorative chrome."""
    s=dk.slide(); dk.header(s, sd["title"])
    y = 1.3
    if sd.get("subtitle"):
        dk.txt(s,Inches(0.55),Inches(y),Inches(12.2),Inches(0.35),[[[sd["subtitle"],{"size":12.5,"italic":True,"color":G700}]]]); y+=0.42
    if sd.get("intro"):
        intro = sd["intro"]
        paras = intro if isinstance(intro,list) and intro and isinstance(intro[0],list) else [[[intro,{"size":12.5,"color":G700}]]] if isinstance(intro,str) else intro
        dk.txt(s,Inches(0.55),Inches(y),Inches(12.2),Inches(0.85),paras,sp_after=4,line_sp=1.15); y+=0.95

    header_row = sd["header"]; rows = sd["rows"]; ncols = len(header_row)
    has_chart = bool(sd.get("chart_image"))
    table_w = Inches(7.7) if has_chart else Inches(12.25)
    nrow = len(rows)+1
    row_h = 0.46
    weights = sd.get("col_widths") or ([2.0]+[1.0]*(ncols-1))
    total_w = sum(weights)

    gt = s.shapes.add_table(nrow, ncols, Inches(0.55), Inches(y), table_w, Inches(row_h*nrow)).table
    for c in range(ncols):
        gt.columns[c].width = int(table_w * (weights[c]/total_w))
    for r in range(nrow):
        gt.rows[r].height = Inches(row_h)
        row = header_row if r==0 else rows[r-1]
        for c in range(ncols):
            cell = row[c] if c < len(row) else ""
            text, tcolor = _cell_text_color(cell)
            tc = gt.cell(r,c); tc.margin_left=Inches(0.12); tc.margin_right=Inches(0.08)
            tc.margin_top=Inches(0.02); tc.margin_bottom=Inches(0.02); tc.vertical_anchor=MSO_ANCHOR.MIDDLE
            p = tc.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            run = p.add_run(); run.text = text; run.font.name=FONT
            run.font.size = Pt(12.5 if r==0 else 12)
            if r==0:
                tc.fill.solid(); tc.fill.fore_color.rgb=GREEN; run.font.bold=True; run.font.color.rgb=WHITE
            else:
                tc.fill.solid(); tc.fill.fore_color.rgb=WHITE if r%2 else G100
                run.font.color.rgb = tcolor
                if c==0: run.font.bold=True
    y_after = y + row_h*nrow + 0.18

    if has_chart:
        img = os.path.join(sd.get("_cdir",""), sd["chart_image"])
        if os.path.exists(img):
            s.shapes.add_picture(img, Inches(8.45), Inches(y), width=Inches(4.3))

    if sd.get("caption"):
        cap = sd["caption"]
        paras = cap if isinstance(cap,list) and cap and isinstance(cap[0],list) else [[[cap,{"size":12,"italic":True,"color":G700}]]]
        dk.txt(s,Inches(0.55),Inches(y_after),Inches(12.2),Inches(0.8),paras,sp_after=4,line_sp=1.15)
        y_after += 0.7

    if sd.get("banner"):
        by = max(y_after+0.1, 6.0)
        dk.rect(s,Inches(0.55),Inches(by),Inches(12.25),Inches(0.85),G100)
        dk.txt(s,Inches(0.8),Inches(by+0.1),Inches(11.7),Inches(0.65),
               [[[sd["banner"].get("label","Bottom line:  "),{"size":13,"bold":True,"color":DARK}],
                 [sd["banner"]["text"],{"size":13,"color":G700}]]],anchor=MSO_ANCHOR.MIDDLE)
    dk.footer(s)

def s_thanks(dk, sd):
    s=dk.slide(); dk.rect(s,0,0,SW,SH,DARK); dk.rect(s,0,0,Inches(0.35),SH,GREEN)
    dk.txt(s,Inches(0.9),Inches(2.7),Inches(11.5),Inches(1.2),[[["Thank you",{"size":44,"bold":True,"color":WHITE}]]])
    dk.txt(s,Inches(0.9),Inches(3.9),Inches(11.5),Inches(1.0),
           [[[sd.get("line","Questions, priorities, or a deeper dive on any section — let's discuss."),{"size":17,"color":CREAM}]]])
    dk.rect(s,Inches(0.9),Inches(4.9),Inches(1.5),Pt(3),GOLD)
    dk.txt(s,Inches(0.9),Inches(5.15),Inches(11),Inches(0.5),[[["Sophie Society  ·  Amazon PPC & Growth",{"size":13,"bold":True,"color":MINT}]]])

R={"cover":s_cover,"exec_summary":s_exec,"kpi_table":s_kpi,"narrative":s_narr,
   "bullets":s_bullets,"two_card":s_twocard,"pl_table":s_pl,"data_table":s_data_table,
   "thankyou":s_thanks}

def main():
    if len(sys.argv)<2:
        print("usage: build_deck.py <report_data.json>"); sys.exit(1)
    jp=os.path.abspath(sys.argv[1])
    with open(jp) as f: data=json.load(f)
    base=os.path.dirname(jp); cdir=os.path.join(base, data.get("charts_dir","charts"))
    dk=Deck(data["meta"]["brand"])
    for sd in data["slides"]:
        t=sd.get("type")
        if t=="chart": s_chart(dk, sd, cdir)
        elif t=="data_table": sd["_cdir"]=cdir; R[t](dk, sd)
        elif t in R: R[t](dk, sd)
        else: print("WARN unknown slide type:", t)
    out=os.path.join(base, data["meta"]["filename"])
    dk.p.save(out); print("SAVED:", out, "| slides:", len(dk.p.slides._sldIdLst))

if __name__=="__main__":
    main()
