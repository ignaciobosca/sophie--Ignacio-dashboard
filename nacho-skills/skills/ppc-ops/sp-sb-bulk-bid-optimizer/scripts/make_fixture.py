"""Build a synthetic Amazon Ads bulk export with SP + SB sheets, shaped like a real one."""
import openpyxl, random

random.seed(7)

SP_HDR = ["Product","Entity","Operation","Campaign ID","Ad Group ID","Portfolio ID","Ad ID",
"Keyword ID","Product Targeting ID","Campaign Name","Ad Group Name",
"Campaign Name (Informational only)","Ad Group Name (Informational only)",
"Portfolio Name (Informational only)","Start Date","End Date","Targeting Type","State",
"Campaign State (Informational only)","Ad Group State (Informational only)",
"Campaign Serving Status (Informational only)","Daily Budget","SKU","ASIN",
"Ad Group Default Bid","Bid","Keyword Text","Native Language Keyword","Native Language Locale",
"Match Type","Bidding Strategy","Placement","Percentage","Product Targeting Expression",
"Resolved Product Targeting Expression","Impressions","Clicks","Click-through Rate","Spend",
"Sales","Orders","Units","Conversion Rate","ACOS","CPC","ROAS"]

SB_HDR = ["Product","Entity","Operation","Campaign ID","Draft Campaign ID","Portfolio ID",
"Ad Group ID","Keyword ID","Product Targeting ID","Campaign Name","Ad Group Name",
"Campaign Name (Informational only)","Ad Group Name (Informational only)",
"Portfolio Name (Informational only)","Start Date","End Date","State",
"Campaign State (Informational only)","Ad Group State (Informational only)","Budget Type",
"Daily Budget","Bid Optimization","Bid Multiplier","Bid","Keyword Text","Match Type",
"Product Targeting Expression","Ad Format","Landing Page URL","Landing Page ASINs","Brand Entity ID",
"Brand Name","Brand Logo Asset ID","Creative Headline","Creative ASINs","Video Media IDs",
"Video Asset IDs","Placement","Percentage","Cost Type","Impressions","Clicks","Spend","Sales",
"Orders","Units","Viewable Impressions","Video First Quartile Views"]

SP_ST_HDR = ["Product","Entity","Campaign ID","Ad Group ID","Portfolio ID","Keyword ID",
"Campaign Name (Informational only)","Ad Group Name (Informational only)","State",
"Campaign State (Informational only)","Ad Group State (Informational only)","Keyword Text",
"Match Type","Customer Search Term","Impressions","Clicks","Spend","Sales","Orders","Units"]

SB_ST_HDR = list(SP_ST_HDR)


def _norm(s):
    return "".join(ch for ch in str(s).lower() if ch.isalnum())


def row(hdr, **kw):
    r = [""] * len(hdr)
    idx = {_norm(h): i for i, h in enumerate(hdr)}
    for k, v in kw.items():
        key = _norm(k)
        if key in idx:
            r[idx[key]] = v
        else:
            raise KeyError(k)
    return r


wb = openpyxl.Workbook()

# ---------------- Sponsored Products ----------------
ws = wb.active
ws.title = "Sponsored Products Campaigns"
ws.append(SP_HDR)

CAMPS = [
    ("100000000000001", "ACME | SP | Exact | Core Terms", 45.0),
    ("100000000000002", "ACME | SP | Broad | Discovery Bucket With A Rather Long Name", 30.0),
    ("100000000000003", "ACME | SP | Auto | Catch-all", 25.0),
]
for cid, cname, budget in CAMPS:
    ws.append(row(SP_HDR, Product="Sponsored Products", Entity="Campaign", Campaign_ID=cid,
                  Campaign_Name=cname, State="enabled",
                  Campaign_State__Informational_only_="enabled", Daily_Budget=budget,
                  Portfolio_Name__Informational_only_="ACME Core"))
    ws.append(row(SP_HDR, Product="Sponsored Products", Entity="Ad Group", Campaign_ID=cid,
                  Ad_Group_ID=cid + "01", Ad_Group_Name="AG1", State="enabled",
                  Campaign_State__Informational_only_="enabled",
                  Ad_Group_State__Informational_only_="enabled", Ad_Group_Default_Bid=0.75,
                  Campaign_Name__Informational_only_=cname,
                  Ad_Group_Name__Informational_only_="AG1"))
    for pl, pct, cl, sa in [("Placement Top", 25, 420, 980.0), ("Placement Product Page", 0, 310, 240.0),
                            ("Placement Rest Of Search", 0, 190, 300.0), ("Placement Amazon Business", 0, 40, 10.0)]:
        ws.append(row(SP_HDR, Product="Sponsored Products", Entity="Bidding Adjustment",
                      Campaign_ID=cid, Campaign_Name__Informational_only_=cname,
                      Campaign_State__Informational_only_="enabled", Placement=pl, Percentage=pct,
                      Clicks=cl, Sales=sa, Spend=cl * 0.8, Impressions=cl * 40, Orders=int(sa / 32)))

KWS = [
    # (campaign idx, text, match, bid, clicks, orders, sales, spend)
    (0, "acme protein powder", "exact", 1.20, 140, 12, 480.0, 168.0),   # branded, high
    (0, "vanilla protein powder", "exact", 1.10, 96, 7, 305.0, 105.0),  # high tier
    (0, "whey isolate 2lb", "exact", 0.95, 41, 0, 0.0, 39.0),           # pause candidate
    (0, "grass fed whey", "exact", 0.85, 6, 1, 44.0, 5.1),              # medium w/ order
    (0, "cheap protein", "exact", 0.60, 5, 0, 0.0, 3.0),                # medium no order
    (0, "keto shake mix", "exact", 0.55, 1, 0, 0.0, 0.55),              # low tier
    (0, "protein for women", "exact", 0.50, 0, 0, 0.0, 0.0),            # zero click
    (1, "protein powder", "broad", 1.05, 220, 9, 360.0, 231.0),         # broad, poor
    (1, "muscle recovery drink mix for athletes", "broad", 0.90, 58, 0, 0.0, 52.2),
    (1, "acme supplements", "broad", 0.95, 33, 4, 190.0, 31.3),         # branded broad
    (1, "shake", "broad", 0.40, 0, 0, 0.0, 0.0),
    (2, "loose-match", "auto", 0.65, 77, 2, 88.0, 50.0),
    (2, "close-match", "auto", 0.0, 45, 3, 150.0, 29.0),                # uses ad group default
]
for ci, text, mt, bid, cl, orders, sales, spend in KWS:
    cid, cname, _ = CAMPS[ci]
    ws.append(row(SP_HDR, Product="Sponsored Products", Entity="Keyword", Campaign_ID=cid,
                  Ad_Group_ID=cid + "01", Keyword_ID=str(abs(hash(text)) % 10**15),
                  Campaign_Name__Informational_only_=cname,
                  Ad_Group_Name__Informational_only_="AG1",
                  Portfolio_Name__Informational_only_="ACME Core",
                  State="enabled", Campaign_State__Informational_only_="enabled",
                  Ad_Group_State__Informational_only_="enabled", Ad_Group_Default_Bid=0.75,
                  Bid=bid, Keyword_Text=text, Match_Type=mt, Clicks=cl, Orders=orders,
                  Sales=sales, Spend=spend, Impressions=cl * 55, Units=orders))

ws.append(row(SP_HDR, Product="Sponsored Products", Entity="Product Targeting",
              Campaign_ID=CAMPS[2][0], Ad_Group_ID=CAMPS[2][0] + "01",
              Product_Targeting_ID="900000000000123",
              Campaign_Name__Informational_only_=CAMPS[2][1],
              Ad_Group_Name__Informational_only_="AG1", State="enabled",
              Campaign_State__Informational_only_="enabled",
              Ad_Group_State__Informational_only_="enabled", Bid=0.70,
              Product_Targeting_Expression='asin="B0XXXXXXXX"', Clicks=63, Orders=0, Sales=0.0,
              Spend=44.1, Impressions=2100))

# ---------------- SP search terms ----------------
ws2 = wb.create_sheet("SP Search Term Report")
ws2.append(SP_ST_HDR)
for kw, st, cl, sp, orders, sales in [
    ("protein powder", "protein powder for dogs", 38, 41.0, 0, 0.0),
    ("protein powder", "free protein powder samples", 22, 26.4, 0, 0.0),
    ("acme protein powder", "acme protein powder vanilla", 60, 66.0, 6, 240.0),
    ("vanilla protein powder", "vanilla protein powder walmart", 9, 21.0, 0, 0.0),
    # A researched core term that has not converted YET -- the exact case the strategic-keep
    # flag exists for. Meets the negation rule on clicks, but killing it burns the research.
    ("protein powder", "grass fed whey protein isolate", 47, 56.4, 0, 0.0),
]:
    ws2.append(row(SP_ST_HDR, Product="Sponsored Products", Entity="Keyword",
                   Campaign_ID=CAMPS[1][0], Ad_Group_ID=CAMPS[1][0] + "01",
                   Campaign_Name__Informational_only_=CAMPS[1][1],
                   Ad_Group_Name__Informational_only_="AG1", State="enabled",
                   Campaign_State__Informational_only_="enabled",
                   Ad_Group_State__Informational_only_="enabled",
                   Keyword_Text=kw, Match_Type="broad", Customer_Search_Term=st,
                   Clicks=cl, Spend=sp, Orders=orders, Sales=sales, Impressions=cl * 30))

# ---------------- Sponsored Brands ----------------
ws3 = wb.create_sheet("SB Multi Ad Group Campaigns")
ws3.append(SB_HDR)
SB_CAMPS = [
    ("200000000000001", "ACME | SBV | Brand Video", "video", "vid-asset-991"),
    ("200000000000002", "ACME | SB | Product Collection", "productCollection", ""),
    ("200000000000003", "ACME | SB | Store Spotlight", "storeSpotlight", ""),
]
for cid, cname, fmt, vid in SB_CAMPS:
    ws3.append(row(SB_HDR, Product="Sponsored Brands", Entity="Campaign", Campaign_ID=cid,
                   Campaign_Name=cname, Campaign_Name__Informational_only_=cname,
                   State="enabled", Campaign_State__Informational_only_="enabled",
                   Daily_Budget=60.0, Ad_Format=fmt, Video_Asset_IDs=vid, Brand_Name="ACME"))
    ws3.append(row(SB_HDR, Product="Sponsored Brands", Entity="Ad Group", Campaign_ID=cid,
                   Ad_Group_ID=cid + "01", Ad_Group_Name="SB AG",
                   Campaign_Name__Informational_only_=cname,
                   Ad_Group_Name__Informational_only_="SB AG", State="enabled",
                   Campaign_State__Informational_only_="enabled",
                   Ad_Group_State__Informational_only_="enabled"))
    for pl, pct, cl, sa in [("Home Page", 0, 210, 260.0), ("Detail Page", -10, 150, 400.0),
                            ("Other Placements", 0, 95, 60.0)]:
        ws3.append(row(SB_HDR, Product="Sponsored Brands", Entity="Bidding Adjustment by Placement",
                       Campaign_ID=cid, Campaign_Name__Informational_only_=cname,
                       Campaign_State__Informational_only_="enabled", State="enabled",
                       Placement=pl, Percentage=pct, Clicks=cl, Sales=sa, Spend=cl * 1.1,
                       Orders=int(sa / 40), Impressions=cl * 60))

SB_KWS = [
    (0, "acme protein", "exact", 1.60, 88, 5, 260.0, 140.8),
    (0, "protein powder video", "phrase", 1.40, 52, 0, 0.0, 72.8),
    # 14 non-branded orders clears SEG_MIN_ORDERS so the video segment can compute a
    # clicks-to-order threshold at all -- without it SB can never propose a pause, correctly.
    (0, "best protein shake", "broad", 1.30, 130, 14, 700.0, 169.0),
    (0, "vegan protein", "exact", 1.20, 2, 0, 0.0, 2.4),
    (1, "protein powder", "phrase", 1.15, 140, 11, 520.0, 161.0),
    (1, "meal replacement shake", "broad", 1.00, 74, 2, 92.0, 74.0),
    (1, "acme store", "exact", 1.25, 30, 4, 210.0, 37.5),
    (2, "acme brand store", "exact", 1.10, 26, 3, 165.0, 28.6),
    (2, "supplement brands", "broad", 0.95, 61, 0, 0.0, 57.95),
    # High-click zero-order SB row so the fixture actually produces an SB PAUSE -- without one,
    # the SB half of the cross-tab selection bug is untestable.
    (0, "matcha whisk cheap", "broad", 1.20, 130, 0, 0.0, 156.0),
]
for ci, text, mt, bid, cl, orders, sales, spend in SB_KWS:
    cid, cname, fmt, vid = SB_CAMPS[ci]
    ws3.append(row(SB_HDR, Product="Sponsored Brands", Entity="Keyword", Campaign_ID=cid,
                   Ad_Group_ID=cid + "01", Keyword_ID=str(abs(hash("sb" + text)) % 10**15),
                   Campaign_Name__Informational_only_=cname,
                   Ad_Group_Name__Informational_only_="SB AG", State="enabled",
                   Campaign_State__Informational_only_="enabled",
                   Ad_Group_State__Informational_only_="enabled", Bid=bid, Keyword_Text=text,
                   Match_Type=mt, Clicks=cl, Orders=orders, Sales=sales, Spend=spend,
                   Impressions=cl * 70))

# legacy duplicate sheet (Amazon really does this)
ws4 = wb.create_sheet("Sponsored Brands Campaigns")
ws4.append(SB_HDR)
for r in list(ws3.iter_rows(min_row=2, values_only=True))[:8]:
    ws4.append(list(r))

ws5 = wb.create_sheet("SB Search Term Report")
ws5.append(SB_ST_HDR)
for kw, st, cl, sp, orders, sales in [
    ("best protein shake", "best protein shake for weight loss reddit", 44, 61.6, 0, 0.0),
    ("supplement brands", "supplement brands cheap", 31, 29.5, 0, 0.0),
    ("best protein shake", "clean protein shake no additives", 38, 53.2, 0, 0.0),
    ("acme protein", "acme protein powder", 40, 64.0, 4, 190.0),
]:
    ws5.append(row(SB_ST_HDR, Product="Sponsored Brands", Entity="Keyword",
                   Campaign_ID=SB_CAMPS[0][0], Ad_Group_ID=SB_CAMPS[0][0] + "01",
                   Campaign_Name__Informational_only_=SB_CAMPS[0][1],
                   Ad_Group_Name__Informational_only_="SB AG", State="enabled",
                   Campaign_State__Informational_only_="enabled",
                   Ad_Group_State__Informational_only_="enabled", Keyword_Text=kw,
                   Match_Type="broad", Customer_Search_Term=st, Clicks=cl, Spend=sp,
                   Orders=orders, Sales=sales, Impressions=cl * 45))

wb.save("/home/claude/bbo/fixture_bulk.xlsx")
print("wrote fixture_bulk.xlsx", wb.sheetnames)
