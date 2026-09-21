"""Extract the scripts from a SKILL.md and run a full acceptance suite."""
import subprocess, json, sys, pathlib, importlib.util, os

SKILL = sys.argv[1] if len(sys.argv) > 1 else "work_v35.md"
# Wipe the scratch dir every run. Reusing it lets a stale artifact from a PREVIOUS version
# satisfy a check the version under test actually fails -- a test that lies is worse than no test.
OUT = pathlib.Path("verify_out")
if OUT.exists():
    import shutil; shutil.rmtree(OUT)
OUT.mkdir()
lines = pathlib.Path(SKILL).read_text().split("\n")
fences = [i for i, l in enumerate(lines) if l.startswith("```")]
names = ["scorer.py", "applier.py", "renderer.py", "prepare_push.py"]
for (s, e), name in zip(zip(fences[0::2], fences[1::2]), names):
    (OUT / name).write_text("\n".join(lines[s + 1:e]) + "\n")

fails, checks = [], 0
def ok(cond, label, detail=""):
    global checks
    checks += 1
    print(("  PASS  " if cond else "  FAIL  ") + label + (("  -- " + detail) if detail else ""))
    if not cond: fails.append(label)

def load(name):
    spec = importlib.util.spec_from_file_location(name, OUT / name)
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m

print("\n[1] syntax + import")
sc = load("scorer.py"); ap = load("applier.py"); rd = load("renderer.py"); pp = load("prepare_push.py")
ok(True, "all four scripts import cleanly")

# ---------------------------------------------------------------- unit: placement caps
print("\n[2] BUG-04 placement dials are real")
def pl_rows(mode_strategy, cur_top, cur_other, prpc_mult):
    """One campaign, top + rest-of-search, both far above campaign RPC so the cap binds."""
    base = {"entity": "bidding adjustment", "campaign state (informational only)": "enabled",
            "campaign id": "C1", "campaign name (informational only)": "C"}
    rows = []
    for pl, cur in [("Placement Top", cur_top), ("Placement Rest Of Search", cur_other)]:
        r = dict(base); r["placement"] = pl; r["percentage"] = cur
        r["clicks"] = 100.0; r["sales"] = 100.0 * prpc_mult
        rows.append(r)
    cfg = dict(sc.STRATEGIES[mode_strategy]); cfg.update(min_cpc=.2, max_cpc=3, target_acos=.3, brand_terms=[])
    camp = {"C1": {"clicks": 400.0, "sales": 400.0}}   # campaign RPC = 1.00
    return {p["placement"]: p for p in sc.decide_placements(rows, cfg, camp)}

# placement RPC = 3.00 vs campaign 1.00 -> ideal = +200pp, so every cap binds
r = pl_rows("ranking", 10, 10, 3.0)
ok(r["Placement Top"]["new_pct"] == 35, "Ranking Push: Top of Search gets +25pp",
   f"10 -> {r['Placement Top']['new_pct']}")
ok(r["Placement Rest Of Search"]["new_pct"] == 25, "Ranking Push: other placements held at +15pp",
   f"10 -> {r['Placement Rest Of Search']['new_pct']}")
r = pl_rows("profit_first", 10, 10, 3.0)
ok(r["Placement Top"]["new_pct"] == 15, "Profit-First: raises by only +5pp",
   f"10 -> {r['Placement Top']['new_pct']}")
r = pl_rows("balanced", 10, 10, 3.0)
ok(r["Placement Top"]["new_pct"] == 25, "Balanced: +15pp default",
   f"10 -> {r['Placement Top']['new_pct']}")
# weak placement -> full down cap on every mode
r = pl_rows("profit_first", 60, 60, 0.1)
ok(r["Placement Top"]["new_pct"] == 10, "Profit-First still cuts at the full -50pp",
   f"60 -> {r['Placement Top']['new_pct']}")

print("\n[3] BUG-05 SP can create a modifier from 0%")
r = pl_rows("balanced", 0, 0, 3.0)
ok("Placement Top" in r and r["Placement Top"]["new_pct"] == 15,
   "0% placement with strong RPC is now proposed", f"0 -> {r.get('Placement Top',{}).get('new_pct')}")
ok(r["Placement Top"]["creates_modifier"] is True, "and is flagged creates_modifier")
r = pl_rows("balanced", 0, 0, 0.1)   # weak placement, ideal floors at 0
ok(r == {}, "a weak 0% placement is still left alone (no 0 -> 0 noise)")

print("\n[4] BUG-07 max-change clamp is now the outer guarantee")
cfg = dict(sc.STRATEGIES["balanced"]); cfg.update(min_cpc=.2, max_cpc=3, target_acos=.3, brand_terms=[])
acct = {"aov": 40.0, "actc": 25.0, "avg_cpc": .9, "acct_clicks": 5000, "acct_orders": 200,
        "cto_branded": 15.0, "branded_orders": 50, "cto_nonbranded": 28.0, "nonbranded_orders": 150}
# The violation is on the UP side: the clamp caps a rise at +15%, then the floor could push
# straight through it. (The down-clamp is a lower bound, so the floor -- which only ever raises
# -- can never breach it.) 0.50 bid, strong RPC, 7-day window so ref CPC = account average.
row = {"entity": "keyword", "state": "enabled", "campaign state (informational only)": "enabled",
       "ad group state (informational only)": "enabled", "campaign id": "C1", "ad group id": "A1",
       "keyword id": "K1", "keyword text": "widget", "match type": "exact",
       "bid": 0.50, "clicks": 120, "orders": 4, "sales": 300.0, "spend": 100.0}
d = sc.decide_bid(row, cfg, acct, {("C1", "A1"): 2.5}, {}, 7)
cap = sc.r2(0.50 * (1 + cfg["max_change_up"]))
ok(d["new_bid"] <= cap + 1e-9, "a rise respects the +15% up cap (v3.4 returned 0.68)",
   f"new_bid = {d['new_bid']}, cap = {cap}")
ok(d["new_bid"] >= cfg["min_cpc"], "and still sits above Min CPC")
# and the floor is not destroyed -- it still lifts the bid off the modelled value
ok(d["new_bid"] > 0.50, "the bid still rises (the floor is respected within the cap)",
   f"{0.50} -> {d['new_bid']}")

print("\n[5] end-to-end across every strategy")
combos = [("balanced", "both"), ("profit_first", "both"), ("launch", "both"),
          ("ranking", "both"), ("cpc_managed", "sp"), ("balanced", "sp"), ("balanced", "sb")]
env = dict(os.environ, SOPHIE_RUN_TS="2026-08-24T10:00:00")
for strat, prod in combos:
    j = OUT / f"res_{strat}_{prod}.json"
    p = subprocess.run([sys.executable, str(OUT / "scorer.py"), "--input", "fixture_bulk.xlsx",
                        "--strategy", strat, "--products", prod, "--target-acos", "30",
                        "--min-cpc", "0.25", "--max-cpc", "2.00", "--brand-terms", "acme",
                        "--window-days", "30", "--output-json", str(j)],
                       capture_output=True, text=True, env=env)
    good = p.returncode == 0 and j.exists() and "error" not in json.loads(j.read_text())
    ok(good, f"scorer runs: {strat} / {prod}", (p.stderr.strip()[-120:] if not good else p.stdout.strip()[:90]))

print("\n[6] approved xlsx is well-formed")
res = OUT / "res_balanced_both.json"
up, rev = OUT / "upload.xlsx", OUT / "reverse.xlsx"
p = subprocess.run([sys.executable, str(OUT / "applier.py"), "--input", "fixture_bulk.xlsx",
                    "--results", str(res), "--upload-xlsx", str(up), "--reverse-xlsx", str(rev)],
                   capture_output=True, text=True, env=env)
ok(p.returncode == 0, "applier runs", p.stderr.strip()[-140:])
import openpyxl
wb = openpyxl.load_workbook(up)
ok("Sponsored Products Campaigns" in wb.sheetnames, "upload has an SP sheet", str(wb.sheetnames))
ok(any(("brands" in n.lower() or n.lower().startswith("sb ")) for n in wb.sheetnames),
   "upload has an SB sheet", str(wb.sheetnames))
prod_col = {}
mixed = []
for sh in wb.sheetnames:
    ws = wb[sh]
    hdr = [str(c.value).lower().strip() if c.value else "" for c in ws[1]]
    pi = hdr.index("product")
    vals = {ws.cell(r, pi + 1).value for r in range(2, ws.max_row + 1)}
    prod_col[sh] = vals
    if len(vals) > 1: mixed.append(sh)
ok(not mixed, "no sheet mixes ad products", str(prod_col))
ids = {"campaign id", "keyword id", "product targeting id", "ad group id"}
bad_id = []
for sh in wb.sheetnames:
    ws = wb[sh]
    hdr = [str(c.value).lower().strip() if c.value else "" for c in ws[1]]
    for i, hname in enumerate(hdr):
        if hname in ids:
            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, i + 1).value
                if v not in (None, "") and not isinstance(v, str): bad_id.append((sh, hname, v))
ok(not bad_id, "every ID column is stored as text", str(bad_id[:3]))

print("\n[7] dashboard renders")
dash = OUT / "dash.html"
p = subprocess.run([sys.executable, str(OUT / "renderer.py"), "--results", str(res),
                    "--source", "fixture_bulk.xlsx", "--input", "fixture_bulk.xlsx",
                    "--output", str(dash)], capture_output=True, text=True, env=env)
ok(dash.exists() and dash.stat().st_size > 50000, "renderer produces a dashboard",
   f"{dash.stat().st_size // 1024} KB" if dash.exists() else p.stderr.strip()[-140:])
html = dash.read_text()
ok("new Date()" not in html.split("</style>")[-1], "no browser-clock leak in the body JS")
ok("±10 percentage points" not in html, "stale placement copy is gone")
ok("1.5× breakeven" not in html, "stale negation copy is gone")

print("\n[8] determinism")
d2 = OUT / "dash2.html"
subprocess.run([sys.executable, str(OUT / "renderer.py"), "--results", str(res), "--source",
                "fixture_bulk.xlsx", "--input", "fixture_bulk.xlsx", "--output", str(d2)],
               capture_output=True, env=env)
ok(dash.read_bytes() == d2.read_bytes(), "same inputs produce a byte-identical dashboard")
r2 = OUT / "res2.json"
subprocess.run([sys.executable, str(OUT / "scorer.py"), "--input", "fixture_bulk.xlsx",
                "--strategy", "balanced", "--products", "both", "--target-acos", "30",
                "--min-cpc", "0.25", "--max-cpc", "2.00", "--brand-terms", "acme",
                "--window-days", "30", "--output-json", str(r2)], capture_output=True, env=env)
ok(res.read_text() == r2.read_text(), "same inputs produce byte-identical scorer JSON")


print("\n[9] v3.6 additions")
# currency flows all the way to the dashboard
gbp = OUT / "res_gbp.json"; gdash = OUT / "dash_gbp.html"
cp = subprocess.run([sys.executable, str(OUT / "scorer.py"), "--input", "fixture_bulk.xlsx",
                "--strategy", "balanced", "--products", "both", "--target-acos", "30",
                "--min-cpc", "0.25", "--max-cpc", "2.00", "--brand-terms", "acme",
                "--currency-symbol", "£", "--window-days", "30",
                "--output-json", str(gbp)], capture_output=True, text=True, env=env)
ok(cp.returncode == 0, "scorer accepts --currency-symbol", cp.stderr.strip()[-100:])
cfgc = (json.loads(gbp.read_text())["config"].get("currency_symbol") if gbp.exists() else None)
ok(cfgc == "£", "scorer persists the currency symbol", f"config.currency_symbol = {cfgc!r}")
subprocess.run([sys.executable, str(OUT / "renderer.py"), "--results", str(gbp), "--source",
                "fixture_bulk.xlsx", "--input", "fixture_bulk.xlsx", "--output", str(gdash)],
               capture_output=True, env=env)
g = gdash.read_text() if gdash.exists() else ""
ok(bool(g) and "£" in g, "dashboard labels money in the chosen currency",
   "" if g else "renderer produced nothing (scorer step failed)")
ok('"currency_symbol": "\\u00a3"' in g or '"currency_symbol": "£"' in g,
   "approved-JSON payload carries it too")

# search-term state gate
st_no_state = {"keyword text": "widget", "customer search term": "widget for dogs",
               "clicks": 60.0, "spend": 90.0, "orders": 0.0, "match type": "broad",
               "campaign name (informational only)": "C", "ad group name (informational only)": "A"}
acct2 = dict(acct)
ok(len(sc.find_negations([st_no_state], {"brand_terms": []}, acct2)) == 1,
   "a search-term row with no State column is still evaluated")
ok(len(sc.find_negations([dict(st_no_state, **{"state": "paused"})], {"brand_terms": []}, acct2)) == 0,
   "an explicitly paused search-term row is still skipped")

# distinct-term count
d = json.loads(res.read_text())
ok("negations_distinct_terms" in d["counts"], "counts carry a distinct-term tally",
   f"{d['counts'].get('negations')} rows / {d['counts'].get('negations_distinct_terms')} terms")

# silent-zero warning when the SP search term sheet is absent
import openpyxl as _ox
nost = OUT / "fixture_no_st.xlsx"
_wb = _ox.load_workbook("fixture_bulk.xlsx"); del _wb["SP Search Term Report"]; _wb.save(nost)
r3 = OUT / "res_nost.json"
subprocess.run([sys.executable, str(OUT / "scorer.py"), "--input", str(nost), "--strategy",
                "balanced", "--products", "sp", "--target-acos", "30", "--min-cpc", "0.25",
                "--max-cpc", "2.00", "--brand-terms", "acme", "--window-days", "30",
                "--output-json", str(r3)], capture_output=True, env=env)
w = " ".join(json.loads(r3.read_text())["warnings"])
ok("negations were skipped, not found to be zero" in w,
   "a missing SP search-term sheet is called out, not reported as zero waste")

print("\n[10] v3.7 strategic-keyword safety flag (Negations only)")
KWF = {"xlsx": "kw_research.xlsx", "csv": "kw_research.csv", "txt": "kw_research.txt"}
shapes = {}
for ext, f in KWF.items():
    if not pathlib.Path(f).exists():
        ok(False, f"keyword research fixture present: {f}"); continue
    j = OUT / f"res_strategic_{ext}.json"
    cp = subprocess.run([sys.executable, str(OUT / "scorer.py"), "--input", "fixture_bulk.xlsx",
                         "--strategy", "balanced", "--products", "both", "--target-acos", "30",
                         "--min-cpc", "0.25", "--max-cpc", "2.00", "--brand-terms", "acme",
                         "--window-days", "30", "--keyword-research", f,
                         "--output-json", str(j)], capture_output=True, text=True, env=env)
    good = cp.returncode == 0 and j.exists()
    ok(good, f"scorer reads a {ext} keyword list", cp.stderr.strip()[-100:] if not good else "")
    if good:
        dd = json.loads(j.read_text())
        shapes[ext] = [(n["customer_search_term"], n.get("strategic"))
                       for n in dd["negations"] + dd["sb"]["negations"]]

ok(len(shapes) == 3 and len(set(map(str, shapes.values()))) == 1,
   "all three file formats produce the identical flagging")

_sfile = OUT / "res_strategic_xlsx.json"
if not _sfile.exists():
    ok(False, "strategic-flag run produced results (skipping the rest of [10])",
       "scorer has no --keyword-research support")
    sd = {"config": {}, "negations": [], "sb": {"negations": []}, "counts": {},
          "changes": None, "placement_changes": None, "budget_changes": None}
else:
    sd = json.loads(_sfile.read_text())
sp_neg, sb_neg = sd["negations"], sd["sb"]["negations"]
ok(sd["config"].get("strategic_terms_loaded") == 5,
   "section headers and numeric noise are not treated as keywords",
   f"loaded {sd['config'].get('strategic_terms_loaded')} of 7 rows")
ok(bool(sp_neg) and sp_neg[0].get("strategic") is True,
   "a researched SP term sorts to the top and is flagged",
   sp_neg[0]["customer_search_term"] if sp_neg else "no rows")
ok(bool(sb_neg) and sb_neg[0].get("strategic") is True,
   "the flag applies to Sponsored Brands too",
   sb_neg[0]["customer_search_term"] if sb_neg else "no rows")
ok(bool(sp_neg) and sp_neg[0]["reason"].startswith("Don't negate"),
   "the flagged row carries the keep instruction, not the negate reason")
ok(all(not n.get("strategic") for n in sp_neg[1:]),
   "unresearched terms are untouched", f"{len(sp_neg) - 1} still negate")
ok(sd["counts"].get("negations_strategic") == 1, "counts break out the strategic keeps")

# case-insensitive: the research file holds "Clean Protein Shake No Additives" in title case
ok(any(n["customer_search_term"] == "clean protein shake no additives" and n.get("strategic")
       for n in sb_neg), "matching is case-insensitive")

# SCOPE: negations only. Bids, pauses, placements and budgets must be byte-identical.
base = json.loads((OUT / "res_balanced_both.json").read_text())
same = sd["changes"] is not None and all(
    json.dumps(base[k], sort_keys=True, default=str) ==
    json.dumps(sd[k], sort_keys=True, default=str)
    for k in ("changes", "placement_changes", "budget_changes"))
ok(same, "supplying keyword research changes NO bid, pause, placement or budget")
ok(json.dumps(base["sb"]["changes"], sort_keys=True, default=str) ==
   json.dumps(sd["sb"].get("changes"), sort_keys=True, default=str),
   "and no Sponsored Brands bid or pause either")

# a file with nothing usable warns rather than silently doing nothing
empty = OUT / "empty_kw.txt"; empty.write_text("\n12345\n\n")
j = OUT / "res_strategic_empty.json"
subprocess.run([sys.executable, str(OUT / "scorer.py"), "--input", "fixture_bulk.xlsx",
                "--strategy", "balanced", "--products", "sp", "--target-acos", "30",
                "--min-cpc", "0.25", "--max-cpc", "2.00", "--brand-terms", "acme",
                "--window-days", "30", "--keyword-research", str(empty),
                "--output-json", str(j)], capture_output=True, env=env)
ok(j.exists() and "yielded no usable terms" in " ".join(json.loads(j.read_text())["warnings"]),
   "an unusable keyword file warns instead of silently skipping")

# dashboard: banner, badge, unticked default
sdash = OUT / "dash_strategic.html"
subprocess.run([sys.executable, str(OUT / "renderer.py"), "--results",
                str(OUT / "res_strategic_xlsx.json"), "--source", "fixture_bulk.xlsx",
                "--input", "fixture_bulk.xlsx", "--output", str(sdash)],
               capture_output=True, env=env)
sh = sdash.read_text() if sdash.exists() else ""
ok("strategic-banner" in sh, "the dashboard shows the strategic banner")
ok("Strategic keep" in sh, "and badges the row")
ok("tr.row-strategic.row-off td{opacity:1}" in sh,
   "the unticked strategic row is not dimmed into illegibility")

print("\n[11] v3.8 selection is scoped to its own tab (reported live)")
import asyncio
try:
    from playwright.async_api import async_playwright
except ImportError:
    ok(False, "playwright available for the selection tests", "pip install playwright")
    async_playwright = None

SEL_DASH = OUT / "dash_sel.html"
subprocess.run([sys.executable, str(OUT / "renderer.py"), "--results", str(res), "--source",
                "fixture_bulk.xlsx", "--input", "fixture_bulk.xlsx", "--output", str(SEL_DASH)],
               capture_output=True, env=env)

async def selection_probe():
    out = {}
    async with async_playwright() as pw:
        b = await pw.chromium.launch()
        pg = await b.new_page(viewport={"width": 1440, "height": 1000})
        await pg.goto(SEL_DASH.resolve().as_uri()); await pg.wait_for_timeout(350)
        async def ticks(panel, tbl):
            # NOTE [type=checkbox] is load-bearing: input.edit-bid carries the same
            # data-tbl/data-idx pair, so without it the editable bid fields get counted as ticks.
            return await pg.evaluate(
                """([p,t])=>[...document.querySelectorAll(`.tab-panel[data-panel='${p}'] input[type=checkbox][data-tbl='${t}'][data-idx]`)].map(c=>c.checked)""",
                [panel, tbl])
        out["bids_n"] = len(await ticks("bids", "bids"))
        out["pauses_n"] = len(await ticks("pauses", "bids"))
        out["sbbids_n"] = len(await ticks("sbbids", "sbbids"))
        out["sbpauses_n"] = len(await ticks("sbpauses", "sbbids"))

        # --- Daian's exact steps: deselect-all on SP Pauses must not touch SP Bids ---
        await pg.click(".tab-btn[data-tab='pauses']"); await pg.wait_for_timeout(120)
        await pg.click(".tab-panel[data-panel='pauses'] input[data-select-all]")
        await pg.wait_for_timeout(120)
        out["after_pause_none__bids"] = await ticks("bids", "bids")
        out["after_pause_none__pauses"] = await ticks("pauses", "bids")

        # --- and the reverse: select-all on SP Bids must not re-enable SP Pauses ---
        await pg.click(".tab-btn[data-tab='bids']"); await pg.wait_for_timeout(120)
        await pg.click(".tab-panel[data-panel='bids'] .mini-btn[data-bulk='all']")
        await pg.wait_for_timeout(120)
        out["after_bids_all__bids"] = await ticks("bids", "bids")
        out["after_bids_all__pauses"] = await ticks("pauses", "bids")

        # --- SB mirror of the same defect ---
        await pg.click(".tab-btn[data-tab='sbpauses']"); await pg.wait_for_timeout(120)
        await pg.click(".tab-panel[data-panel='sbpauses'] .mini-btn[data-bulk='none']")
        await pg.wait_for_timeout(120)
        out["after_sbpause_none__sbbids"] = await ticks("sbbids", "sbbids")
        out["after_sbpause_none__sbpauses"] = await ticks("sbpauses", "sbbids")

        # --- header checkbox tells the truth about its own panel ---
        await pg.click(".tab-btn[data-tab='bids']"); await pg.wait_for_timeout(120)
        await pg.click(".tab-panel[data-panel='bids'] tbody tr:first-child input[type=checkbox][data-idx]")
        await pg.wait_for_timeout(120)
        out["head_after_one_untick"] = await pg.evaluate(
            """()=>{const h=document.querySelector(".tab-panel[data-panel='bids'] input[data-select-all]");
                    return {checked:h.checked, indeterminate:h.indeterminate};}""")
        await b.close()
    return out

if async_playwright:
    sel = asyncio.run(selection_probe())
    ok(sel["bids_n"] > 0 and sel["pauses_n"] > 0,
       "fixture has both SP bids and SP pauses to cross-test",
       f"{sel['bids_n']} bids / {sel['pauses_n']} pauses")
    ok(sel["sbbids_n"] > 0 and sel["sbpauses_n"] > 0,
       "and both SB bids and SB pauses",
       f"{sel['sbbids_n']} bids / {sel['sbpauses_n']} pauses")
    ok(all(sel["after_pause_none__bids"]),
       "deselect-all on SP Pauses leaves SP Bids untouched  <-- the reported bug",
       f"{sum(sel['after_pause_none__bids'])}/{len(sel['after_pause_none__bids'])} still ticked")
    ok(not any(sel["after_pause_none__pauses"]),
       "...and does deselect SP Pauses itself")
    ok(not any(sel["after_bids_all__pauses"]),
       "select-all on SP Bids does not re-enable SP Pauses  <-- the reverse case")
    ok(all(sel["after_bids_all__bids"]), "...and does select SP Bids itself")
    ok(all(sel["after_sbpause_none__sbbids"]),
       "the same holds for SB Bids vs SB Pauses",
       f"{sum(sel['after_sbpause_none__sbbids'])}/{len(sel['after_sbpause_none__sbbids'])} still ticked")
    ok(not any(sel["after_sbpause_none__sbpauses"]), "...and SB Pauses did clear")
    h = sel["head_after_one_untick"]
    ok(h["indeterminate"] is True and h["checked"] is False,
       "header checkbox goes indeterminate on a partial selection", str(h))

print("\n[12] no bulk control anywhere leaks into another tab (exhaustive)")
# [11] pins the two cases that were reported. This pins the general INVARIANT: for every panel,
# every bulk control in it may change that panel's rows and NOTHING else. It covers tabs nobody
# has complained about, and it fails the day a new tab ships with a document-wide handler.
# One page load; state is reset in-page between controls rather than by reloading, which keeps
# ~30 control checks inside a couple of seconds.

async def leak_matrix():
    findings = []
    async with async_playwright() as pw:
        b = await pw.chromium.launch()
        pg = await b.new_page(viewport={"width": 1440, "height": 1000})
        await pg.goto(SEL_DASH.resolve().as_uri()); await pg.wait_for_timeout(300)

        RESET = """()=>{try{localStorage.clear()}catch(e){}
          document.querySelectorAll("input[type=checkbox][data-tbl][data-idx]").forEach(c=>{
            c.checked=true; const tr=c.closest('tr'); if(tr) tr.classList.remove('row-off');});
          document.querySelectorAll('input[data-select-all]').forEach(h=>{h.checked=true;h.indeterminate=false;});}"""
        SNAP = """()=>{const out={};
          document.querySelectorAll('.tab-panel').forEach(p=>{
            out[p.dataset.panel]=[...p.querySelectorAll("input[type=checkbox][data-tbl][data-idx]")]
              .map(c=>c.dataset.tbl+':'+c.dataset.idx+'='+(c.checked?1:0)).join(',');});
          return out;}"""

        plan = await pg.evaluate("""()=>{
          const out=[];
          document.querySelectorAll('.tab-panel').forEach(p=>{
            const panel=p.dataset.panel;
            if(!p.querySelector("input[type=checkbox][data-tbl][data-idx]")) return;
            if(p.querySelector('input[data-select-all]')) out.push([panel,'selectall',null]);
            p.querySelectorAll('.mini-btn[data-bulk]').forEach(b=>out.push([panel,'bulk',b.dataset.bulk]));
          });
          return out;}""")

        for panel, kind, mode in plan:
            # Open the tab first -- an inactive panel is display:none and cannot be clicked,
            # and opening it is what an operator does anyway.
            await pg.click(f".tab-btn[data-tab='{panel}']")
            await pg.evaluate(RESET)
            before = await pg.evaluate(SNAP)
            sel = (f".tab-panel[data-panel='{panel}'] input[data-select-all]" if kind == 'selectall'
                   else f".tab-panel[data-panel='{panel}'] .mini-btn[data-bulk='{mode}']")
            try:
                await pg.click(sel)
            except Exception as exc:
                findings.append((panel, kind, mode, f"click failed: {type(exc).__name__}"))
                continue
            after = await pg.evaluate(SNAP)
            for other in before:
                if other != panel and before[other] != after[other]:
                    findings.append((panel, kind, mode, f"also changed '{other}'"))
        await b.close()
    return plan, findings

if async_playwright:
    plan, leaks = asyncio.run(leak_matrix())
    ok(len(plan) >= 20, "every panel's bulk controls were exercised",
       f"{len(plan)} controls across {len({p for p,_,_ in plan})} panels")
    ok(not leaks, "no control changes a tab other than its own",
       "; ".join(f"{p}/{k}{'/'+str(m) if m else ''} -> {why}" for p, k, m, why in leaks[:4]) or "clean")

# localStorage round-trip: the saved ticks must come back per-tab, not smeared across tabs
async def persist_probe():
    async with async_playwright() as pw:
        b = await pw.chromium.launch()
        pg = await b.new_page(viewport={"width": 1440, "height": 1000})
        await pg.goto(SEL_DASH.resolve().as_uri()); await pg.wait_for_timeout(300)
        await pg.click(".tab-btn[data-tab='pauses']"); await pg.wait_for_timeout(100)
        await pg.click(".tab-panel[data-panel='pauses'] .mini-btn[data-bulk='none']")
        await pg.wait_for_timeout(150)
        await pg.reload(); await pg.wait_for_timeout(400)
        out = await pg.evaluate("""()=>{
          const g=(p,t)=>[...document.querySelectorAll(`.tab-panel[data-panel='${p}'] input[type=checkbox][data-tbl='${t}'][data-idx]`)].map(c=>c.checked);
          return {bids:g('bids','bids'), pauses:g('pauses','bids')};}""")
        await b.close()
    return out

if async_playwright:
    pr = asyncio.run(persist_probe())
    ok(all(pr["bids"]) and not any(pr["pauses"]),
       "saved state reloads per-tab too (Pauses off, Bids intact)",
       f"bids {sum(pr['bids'])}/{len(pr['bids'])} on, pauses {sum(pr['pauses'])}/{len(pr['pauses'])} on")

print("\n[13] the three output routes agree (found by an operator run, not by this suite)")
# An operator downloads the BROWSER xlsx; the Python applier is a second route; the approved JSON
# is a third that writes to the live account. All three must say the same thing about the same
# row. On Profit-First they did not: split-policy rows were paused / floored / left below Min CPC
# depending on which route you took. Profit-First is the only strategy with below_min='split',
# which is why every earlier version of this suite missed it.

async def route_compare(strategy):
    jf = OUT / f"routes_{strategy}.json"
    subprocess.run([sys.executable, str(OUT / "scorer.py"), "--input", "fixture_bulk.xlsx",
                    # Min CPC is set deliberately HIGH so the fixture actually produces below-min
                    # rows on BOTH ad products. At a lower floor the SB side has none, and the
                    # v3.8 push/xlsx divergence -- which only ever showed on SB -- hides.
                    "--strategy", strategy, "--products", "both", "--target-acos", "30",
                    "--min-cpc", "0.75", "--max-cpc", "2.00", "--brand-terms", "acme",
                    "--window-days", "30", "--output-json", str(jf)], capture_output=True, env=env)
    dash = OUT / f"routes_{strategy}.html"
    subprocess.run([sys.executable, str(OUT / "renderer.py"), "--results", str(jf), "--source",
                    "fixture_bulk.xlsx", "--input", "fixture_bulk.xlsx", "--output", str(dash)],
                   capture_output=True, env=env)
    pyx = OUT / f"routes_{strategy}_py.xlsx"
    subprocess.run([sys.executable, str(OUT / "applier.py"), "--input", "fixture_bulk.xlsx",
                    "--results", str(jf), "--upload-xlsx", str(pyx),
                    "--reverse-xlsx", str(OUT / f"routes_{strategy}_rev.xlsx")],
                   capture_output=True, env=env)
    async with async_playwright() as pw:
        b = await pw.chromium.launch()
        ctx = await b.new_context(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
        pg = await ctx.new_page()
        await pg.goto(dash.resolve().as_uri()); await pg.wait_for_timeout(300)
        # Visit whatever tabs this build actually renders. A hardcoded list silently stops
        # unlocking the download gate the moment a new tab is added (v4.0 added two).
        for t in await pg.evaluate("()=>[...document.querySelectorAll('.tab-btn')].map(b=>b.dataset.tab)"):
            try: await pg.click(f".tab-btn[data-tab='{t}']")
            except Exception: pass
        await pg.wait_for_timeout(120)
        await pg.click(".tab-btn[data-tab='summary']")
        async with pg.expect_download() as di: await pg.click("#btn-forward")
        brx = OUT / f"routes_{strategy}_browser.xlsx"
        await (await di.value).save_as(brx)
        async with pg.expect_download() as di2: await pg.click("#btn-approved-json")
        bjs = OUT / f"routes_{strategy}_approved.json"
        await (await di2.value).save_as(bjs)
        has_sb_bm = await pg.evaluate("()=>document.querySelectorAll(\"select.bm-override[data-tbl='sbbids']\").length")
        n_sb_belowmin = len([c for c in json.loads(jf.read_text()).get("sb", {}).get("changes", [])
                             if c.get("flag_below_min")])
        await b.close()
    return jf, pyx, brx, bjs, has_sb_bm, n_sb_belowmin

def sheet_rows(path):
    wb = openpyxl.load_workbook(path); out = {}
    for sh in wb.sheetnames:
        ws = wb[sh]; hdr = [str(c.value).lower().strip() if c.value else "" for c in ws[1]]
        def col(n): return hdr.index(n) if n in hdr else None
        ei, bi, si, pci, pli = col("entity"), col("bid"), col("state"), col("percentage"), col("placement")
        ki, pti, kid, ptid = col("keyword text"), col("product targeting expression"), col("keyword id"), col("product targeting id")
        for r in range(2, ws.max_row + 1):
            g = lambda i: (ws.cell(r, i + 1).value if i is not None else None)
            key = (sh, str(g(ei)).lower(), str(g(kid) or g(ptid) or ""), str(g(pli) or ""))
            out[key] = {"bid": g(bi), "state": g(si), "pct": g(pci), "entity": g(ei),
                        "text": g(ki) or g(pti)}
    return out

if async_playwright:
    for strat in ("profit_first", "balanced"):
        jf, pyx, brx, bjs, has_sb_bm, n_sb_bm = asyncio.run(route_compare(strat))
        P, B = sheet_rows(pyx), sheet_rows(brx)
        only_p, only_b = set(P) - set(B), set(B) - set(P)
        diffs = [(k, P[k]["bid"], P[k]["state"], B[k]["bid"], B[k]["state"])
                 for k in set(P) & set(B)
                 if (str(P[k]["bid"]), str(P[k]["state"])) != (str(B[k]["bid"]), str(B[k]["state"]))]
        ok(not only_p and not only_b,
           f"{strat}: python applier and browser writer emit the same rows",
           f"py-only {len(only_p)}, browser-only {len(only_b)}")
        ok(not diffs, f"{strat}: ...with the same bid and state on every row",
           "; ".join(f"{k[2] or k[3]}: py={a}/{b} browser={c}/{d}" for k, a, b, c, d in diffs[:3]) or "identical")

        # The push payload must say the same thing as the xlsx for every target. An absolute
        # "never below Min CPC" assertion would be wrong: under 'split' (with orders) and 'keep'
        # a below-min bid is the operator's chosen outcome and all three routes carry it. What
        # must never happen is the routes DISAGREEING -- that is what shipped in v3.8, where the
        # push sent raw modelled bids while the xlsx floored or paused them.
        doc = json.loads(bjs.read_text())
        cfg_j = json.loads(jf.read_text())["config"]
        floor = cfg_j["min_cpc"]; sb_floor = (cfg_j.get("sb") or {}).get("min_cpc", floor)
        policy = cfg_j.get("below_min")
        xlsx_bid = {}
        for k, v in B.items():
            if v["bid"] not in (None, "") and str(v["entity"] or "").lower() in ("keyword", "product targeting"):
                xlsx_bid[(k[0], k[2])] = float(v["bid"])
        mism = []
        for arr, sheet_hint in ((doc["target_bids"], "sponsored products"), (doc["sb_target_bids"], "sb ")):
            for x in arr:
                hit = [vv for (sh, tid), vv in xlsx_bid.items()
                       if tid == str(x["targetId"]) and sheet_hint.split()[0] in sh.lower()]
                if hit and abs(hit[0] - x["final_bid"]) > 1e-9:
                    mism.append((x.get("text") or x["targetId"], hit[0], x["final_bid"]))
        ok(not mism, f"{strat}: the push payload matches the approved xlsx bid for bid  <-- v3.8 defect",
           "; ".join(f"{t}: xlsx={a} push={b}" for t, a, b in mism[:3]) or "identical")
        # and where the policy is floor/pause, nothing may sit under the floor at all
        if policy in ("floor", "pause"):
            under = [x for x in doc["target_bids"] if x["final_bid"] < floor]
            sb_under = [x for x in doc["sb_target_bids"] if x["final_bid"] < sb_floor]
            ok(not under and not sb_under,
               f"{strat}: under a '{policy}' policy nothing is pushed below Min CPC",
               f"SP {len(under)}, SB {len(sb_under)}")

        # placement Entity must match what Amazon uses on each sheet
        ents = {}
        for k, v in B.items():
            if str(v["entity"] or "").lower().startswith("bidding"):
                ents.setdefault(k[0], set()).add(v["entity"])
        sp_ok = all(e == "Bidding Adjustment" for s_, es in ents.items()
                    if "sponsored products" in s_.lower() for e in es)
        sb_ok = all(e == "Bidding Adjustment by Placement" for s_, es in ents.items()
                    if "sponsored products" not in s_.lower() for e in es)
        ok(sp_ok and sb_ok, f"{strat}: placement Entity matches Amazon's per-sheet name", str(ents) or "no placement rows")

        if n_sb_bm:
            ok(has_sb_bm == n_sb_bm,
               f"{strat}: every below-min SB row gets a BM override control",
               f"{has_sb_bm} controls for {n_sb_bm} flagged rows")

print("\n[14] CPC-Managed: a suppressed check never looks like a passed one")
def felix_run(mode, yesterday=False):
    j = OUT / f"felix_{mode}{'_y' if yesterday else ''}.json"
    cmd = [sys.executable, str(OUT / "scorer.py"), "--input", "fixture_bulk.xlsx",
           "--strategy", "cpc_managed", "--felix-mode", mode, "--products", "both",
           "--target-acos", "30", "--min-cpc", "0.25", "--max-cpc", "2.00",
           "--brand-terms", "acme", "--window-days", "30", "--output-json", str(j)]
    if yesterday: cmd += ["--yesterday-input", "fixture_bulk.xlsx"]
    subprocess.run(cmd, capture_output=True, env=env)
    return json.loads(j.read_text()) if j.exists() else {}

d = felix_run("daily")
w = " ".join(d.get("warnings", []))
ok(d.get("counts", {}).get("budget_changes") == 0 and "SKIPPED, not passed" in w,
   "daily with no Yesterday bulk says the budget check was skipped", "warned" if "SKIPPED" in w else "SILENT ZERO")

dy = felix_run("daily", yesterday=True)
wy = " ".join(dy.get("warnings", []))
ok("SKIPPED, not passed" not in wy,
   "...and does not cry wolf when the Yesterday file IS supplied")

for mode, absent in (("weekly", "bids and pauses were NOT scored"),
                     ("monthly", "bids and pauses were NOT scored")):
    dm = felix_run(mode); wm = " ".join(dm.get("warnings", []))
    ok(absent in wm, f"{mode} mode explains its empty Bids/Pauses tabs",
       "explained" if absent in wm else "silent")
dw = felix_run("weekly")
ok(dw.get("counts", {}).get("bid_changes") == 0 and dw.get("counts", {}).get("placement_changes", 0) >= 0,
   "weekly scores placements and not bids (the documented separation holds)",
   f"bids={dw.get('counts',{}).get('bid_changes')} placements={dw.get('counts',{}).get('placement_changes')}")
dmo = felix_run("monthly"); wmo = " ".join(dmo.get("warnings", []))
ok("does not touch placements" in wmo, "monthly explains its empty Placements tabs")

print("\n[15] v4.0 quality pass (from a POD-leader review of a live run)")

# --- the upload bug: a Shopper Cohort row must never become a placement row ---
cohort = {"entity": "bidding adjustment", "campaign state (informational only)": "enabled",
          "campaign id": "C1", "campaign name (informational only)": "C",
          "placement": None, "percentage": None, "clicks": 300.0, "sales": 900.0,
          "audience id": "412978094962770620", "shopper cohort type": "Audience Segment",
          "shopper cohort percentage": 50.0}
cfgp = dict(sc.STRATEGIES["balanced"]); cfgp.update(min_cpc=.5, max_cpc=3, target_acos=.45, brand_terms=[])
ok(sc.decide_placements([cohort], cfgp, {"C1": {"clicks": 1000.0, "sales": 1000.0}}) == [],
   "a Shopper Cohort row is never proposed as a placement  <-- the failed-upload bug")
bare = dict(cohort); [bare.pop(k) for k in ("audience id", "shopper cohort type", "shopper cohort percentage")]
ok(sc.decide_placements([bare], cfgp, {"C1": {"clicks": 1000.0, "sales": 1000.0}}) == [],
   "...and neither is any adjustment row with no placement label")
sb_bare = dict(bare); sb_bare["entity"] = "bidding adjustment by placement"
sbcfg = dict(cfgp)
ok(sc.decide_sb_placements([sb_bare], sbcfg, {}) == [], "the same guard applies to Sponsored Brands")

# --- low-data rows are never bid DOWN ---
acct_l = {"aov": 40.0, "actc": 25.0, "avg_cpc": .9, "acct_clicks": 5000, "acct_orders": 200,
          "cto_branded": 15.0, "branded_orders": 50, "cto_nonbranded": 28.0, "nonbranded_orders": 150}
def low_row(clicks, bid=1.00):
    return {"entity": "keyword", "state": "enabled", "campaign state (informational only)": "enabled",
            "ad group state (informational only)": "enabled", "campaign id": "C1", "ad group id": "A1",
            "keyword id": "K1", "keyword text": "widget", "match type": "exact",
            "bid": bid, "clicks": clicks, "orders": 0, "sales": 0.0, "spend": clicks * 0.9}
# ad group RPC well BELOW the bid, so the old rule would have cut
ag_low = {("C1", "A1"): 0.50}
for strat in ("balanced", "launch", "ranking", "profit_first", "cpc_managed"):
    cfl = dict(sc.STRATEGIES[strat]); cfl.update(min_cpc=.2, max_cpc=3, target_acos=.30, brand_terms=[])
    cuts = []
    for n in (1, 2):
        d = sc.decide_bid(low_row(n), cfl, acct_l, ag_low, {}, 30)
        if d and d.get("action") == "bid" and d["new_bid"] < 1.00:
            cuts.append((n, d["new_bid"], d.get("formula", "")[:40]))
    ok(not cuts, f"{strat}: a 1-2 click target is never bid down  <-- reported defect",
       str(cuts[:2]) if cuts else "no cuts")
# and it still steps UP where the strategy says raise
ag_high = {("C1", "A1"): 8.00}
d = sc.decide_bid(low_row(2), dict(sc.STRATEGIES["balanced"], min_cpc=.2, max_cpc=9,
                                   target_acos=.30, brand_terms=[]), acct_l, ag_high, {}, 30)
ok(d and d["new_bid"] > 1.00 and d.get("low_data") is True,
   "Balanced still steps a low-data target UP toward click value, and tags it low_data",
   f"1.00 -> {d['new_bid'] if d else None}")
d = sc.decide_bid(low_row(2), dict(sc.STRATEGIES["profit_first"], min_cpc=.2, max_cpc=9,
                                   target_acos=.30, brand_terms=[]), acct_l, ag_high, {}, 30)
ok(d is None, "Profit-First holds a low-data target rather than buying unproven traffic")

# --- dashboard: low-data tab, sort/filter, dots, save plan ---
async def v40_probe():
    async with async_playwright() as pw:
        b = await pw.chromium.launch()
        ctx = await b.new_context(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
        pg = await ctx.new_page(); errs = []
        pg.on("pageerror", lambda e: errs.append(str(e)[:120]))
        await pg.goto(SEL_DASH.resolve().as_uri()); await pg.wait_for_timeout(350)
        out = {"errs": errs}
        out["tabs"] = await pg.evaluate("()=>[...document.querySelectorAll('.tab-btn')].map(b=>b.dataset.tab)")
        out["dots_before"] = await pg.evaluate("()=>document.querySelectorAll('.tab-btn:not(.seen) .dot').length")
        out["low_in_bids"] = await pg.evaluate("()=>document.querySelectorAll(\"#tbody-bids tr[data-low='1']\").length")
        await pg.click(".tab-btn[data-tab='bids']"); await pg.wait_for_timeout(120)
        base = await pg.evaluate("()=>sheetsFor(false).map(s=>s.rows.length)")
        # An older build has no sort/filter controls -- record that rather than timing out.
        out["has_tools"] = await pg.evaluate("()=>!!document.querySelector(\"select[data-sort='bids']\")")
        if out["has_tools"]:
            await pg.select_option("select[data-sort='bids']", "conf")
            await pg.select_option("select[data-filter='bids']", "High"); await pg.wait_for_timeout(150)
        out["confs"] = await pg.evaluate("()=>[...document.querySelectorAll('#tbody-bids tr[data-idx]')].map(r=>r.dataset.conf)")
        out["export_same"] = (await pg.evaluate("()=>sheetsFor(false).map(s=>s.rows.length)")) == base
        for t in out["tabs"]:
            await pg.click(f".tab-btn[data-tab='{t}']"); await pg.wait_for_timeout(30)
        out["dots_after"] = await pg.evaluate("()=>document.querySelectorAll('.tab-btn:not(.seen) .dot').length")
        await pg.click(".tab-btn[data-tab='summary']"); await pg.wait_for_timeout(100)
        if await pg.evaluate("()=>!!document.getElementById('btn-save-plan')"):
            async with pg.expect_download() as d: await pg.click("#btn-save-plan")
            pl = OUT / "review_plan.json"; await (await d.value).save_as(pl)
            out["plan"] = json.loads(pl.read_text())
        else:
            out["plan"] = {}
        await b.close()
    return out

if async_playwright:
    v = asyncio.run(v40_probe())
    ok(not v["errs"], "no JS errors on the v4.0 dashboard", str(v["errs"][:2]))
    ok("lowdata" in v["tabs"], "a low-data tab is rendered when the split is on", str(v["tabs"]))
    ok(v["low_in_bids"] == 0, "and no low-data row is left in the main Bids tab")
    ok(v.get("has_tools") and
       v["confs"] == sorted(v["confs"], key=lambda x: {"High": 0, "Medium": 1, "Low": 2}.get(x, 9)),
       "sorting by confidence actually orders the rows",
       str(v["confs"][:6]) if v.get("has_tools") else "no sort control on this build")
    ok(v["export_same"], "sorting and filtering NEVER change what gets exported")
    ok(v["dots_before"] > 0 and v["dots_after"] == 0,
       "unreviewed tabs carry a dot that clears once opened",
       f"{v['dots_before']} -> {v['dots_after']}")
    ok(v["plan"].get("schema", "").startswith("sophie.bulk-bid-optimizer.review-plan")
       and "state" in v["plan"],
       "a review plan can be saved for resuming later", v["plan"].get("schema", "?"))

print(f"\n{'='*62}\n{checks - len(fails)}/{checks} checks passed")
if fails:
    print("FAILED:"); [print("  -", f) for f in fails]; sys.exit(1)
print("ALL GREEN")
