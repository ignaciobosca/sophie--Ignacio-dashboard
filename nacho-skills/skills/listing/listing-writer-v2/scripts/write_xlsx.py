#!/usr/bin/env python3
"""
write_xlsx.py — brand-agnostic 3-tab listing workbook generator for listing-writer-v2.

Reads a draft JSON (schema below) and writes the workbook. No product brand is
hardcoded; everything comes from the draft. House style comes from config.py.

Draft JSON schema (all strings authored upstream by the skill, already audited):
{
  "mode": "renovation" | "new",
  "brand": "Acme Goods",
  "product_slug": "learning-tablet",
  "copy": {
     "Brand":        {"value": "...", "limit": "-",        "changed": "..."},
     "Product Title":{"value": "...", "limit": "200 chars","changed": "..."},
     "Bullet 1":     {"value": "...", "limit": "500 chars","changed": "..."},
     ... Bullet 2-5 ...,
     "Product Description": {"value":"...","limit":"2000 chars","changed":"..."},
     "Backend Search Terms":{"value":"...","limit":"250 bytes","changed":"..."}
  },
  "attributes": [
     {"attribute":"Target Age Range","value":"1 to 7 years","source":"seller-confirmed",
      "confidence":"High","action":""},
     ...
  ],
  "images": [
     {"slot":"Hero","goal":"...","cosmo":"Product Type","question":"...",
      "ai_prompt":"...","notes":"..."},
     ... 6 more ...
  ],
  "coverage": {
     "cosmo": [["Product Type","Covered","Title, Hero"], ...],
     "keywords": [["toddler flash cards","Tier 1","Title","head term, vol 5019"], ...]
  },
  "claims_to_verify": ["...", "..."]
}
"""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    from config import STYLE
except Exception:
    STYLE = {"header_bg":"0A2333","secondary_bg":"13835B","flag_color":"EA6C34",
             "row_fill":"F3F3F1","body_color":"2B2B2B","font":"DM Sans"}

def byte_count(s): return len(s.encode("utf-8"))

def build(draft_path, out_dir):
    d = json.load(open(draft_path, encoding="utf-8"))
    brand = d.get("brand","").strip()
    slug  = d.get("product_slug","product").strip().replace(" ","-")
    mode  = d.get("mode","new")
    from datetime import date
    tag = "Renovation" if mode=="renovation" else "Draft"
    out = os.path.join(out_dir, f"Listing_{tag}_{slug}_{date.today().isoformat()}.xlsx")

    F=STYLE
    hdr  = Font(name=F["font"], bold=True, color="FFFFFF", size=11)
    body = Font(name=F["font"], color=F["body_color"], size=10)
    bold = Font(name=F["font"], bold=True, color=F["header_bg"], size=10)
    flag = Font(name=F["font"], bold=True, color=F["flag_color"], size=10)
    navy = PatternFill("solid", fgColor=F["header_bg"])
    grn  = PatternFill("solid", fgColor=F["secondary_bg"])
    orn  = PatternFill("solid", fgColor=F["flag_color"])
    zeb  = PatternFill("solid", fgColor=F["row_fill"])
    wrap = Alignment(wrap_text=True, vertical="top")
    edge = Side(style="thin", color="D8D8D2"); bdr = Border(edge,edge,edge,edge)

    def header_row(ws, fill=navy):
        for c in ws[1]:
            c.font=hdr; c.fill=fill; c.border=bdr
            c.alignment=Alignment(vertical="center", wrap_text=True)

    wb = Workbook()

    # ---------- TAB 1: Listing Copy ----------
    ws = wb.active; ws.title = "Listing Copy"
    ws.append(["Field","Value","Count","TOS Limit","What Changed","Notes"])
    header_row(ws)
    order = ["Brand","Product Title","Item Highlights","Bullet 1","Bullet 2","Bullet 3",
             "Bullet 4","Bullet 5","Product Description","Backend Search Terms"]
    for f in order:
        e = d["copy"].get(f)
        if not e: continue
        v = e["value"]
        if "byte" in e.get("limit",""):
            cnt = f"{byte_count(v)} bytes"
        elif v and e.get("limit","-") != "-":
            cnt = f"{len(v)} chars"
        else:
            cnt = "-"
        changed = e.get("changed","New" if mode=="new" else "")
        ws.append([f, v, cnt, e.get("limit","-"), changed, e.get("notes","")])
    widths = {"A":20,"B":80,"C":12,"D":12,"E":42,"F":34}
    for col,w in widths.items(): ws.column_dimensions[col].width=w
    for i,row in enumerate(ws.iter_rows(min_row=2), start=2):
        for c in row: c.font=body; c.alignment=wrap; c.border=bdr
        ws.cell(i,1).font=bold
        if i%2==0:
            for c in row: c.fill=zeb

    # ---------- TAB 2: Structured Attributes (+ paste block) ----------
    ws2 = wb.create_sheet("Structured Attributes")
    ws2.append(["Attribute","Proposed Value","Source","Confidence","Seller Action"])
    header_row(ws2, grn)
    for a in d.get("attributes",[]):
        ws2.append([a["attribute"], a.get("value",""), a.get("source",""),
                    a.get("confidence",""), a.get("action","")])
    for col,w in zip("ABCDE",[26,52,20,12,52]): ws2.column_dimensions[col].width=w
    for i,row in enumerate(ws2.iter_rows(min_row=2), start=2):
        for c in row: c.font=body; c.alignment=wrap; c.border=bdr
        ws2.cell(i,1).font=bold
        act = (ws2.cell(i,5).value or "")
        if "CONFIRM" in act.upper() or "confirm" in act.lower():
            ws2.cell(i,5).font = flag
        if i%2==0:
            for c in row: c.fill=zeb

    # paste block (clean two-column, for revamp loop + Seller Central paste)
    pb = ws2.max_row + 2
    ws2.cell(pb,1,"PASTE BLOCK (Attribute -> Value)").font=Font(name=F["font"],bold=True,color="FFFFFF",size=11)
    for col in range(1,6): ws2.cell(pb,col).fill=navy
    ws2.append(["Attribute","Value","","",""])
    hr=pb+1
    for c in ws2[hr]: c.font=hdr; c.fill=grn; c.border=bdr
    for a in d.get("attributes",[]):
        if a.get("value") and "CONFIRM" not in (a.get("action","")).upper():
            ws2.append([a["attribute"], a["value"], "","",""])
    for i in range(hr+1, ws2.max_row+1):
        ws2.cell(i,1).font=bold; ws2.cell(i,2).font=body
        for j in (1,2): ws2.cell(i,j).alignment=wrap; ws2.cell(i,j).border=bdr

    # ---------- TAB 3: Image Brief (+ AI prompt) ----------
    ws3 = wb.create_sheet("Image Brief")
    ws3.append(["Slot","Visual Goal","COSMO Dimension","Shopper Question Answered",
                "AI Image Prompt","Production Notes"])
    header_row(ws3)
    for im in d.get("images",[]):
        ws3.append([im.get("slot",""), im.get("goal",""), im.get("cosmo",""),
                    im.get("question",""), im.get("ai_prompt",""), im.get("notes","")])
    for col,w in zip("ABCDEF",[16,40,20,28,60,40]): ws3.column_dimensions[col].width=w
    for i,row in enumerate(ws3.iter_rows(min_row=2), start=2):
        if i-1 > len(d.get("images",[])): break
        for c in row: c.font=body; c.alignment=wrap; c.border=bdr
        ws3.cell(i,1).font=bold
        if i%2==0:
            for c in row: c.fill=zeb

    # Coverage map: COSMO
    cm = ws3.max_row + 2
    ws3.cell(cm,1,"COVERAGE MAP — COSMO").font=Font(name=F["font"],bold=True,color="FFFFFF",size=11)
    for col in range(1,7): ws3.cell(cm,col).fill=navy
    ws3.append(["COSMO Dimension","Status","Placement","","",""])
    hr=cm+1
    for c in ws3[hr]: c.font=hdr; c.fill=grn; c.border=bdr
    for rowv in d.get("coverage",{}).get("cosmo",[]):
        ws3.append(list(rowv)+["","",""][:6-len(rowv)])
    for i in range(hr+1, ws3.max_row+1):
        ws3.cell(i,1).font=bold
        for j in range(1,4):
            ws3.cell(i,j).alignment=wrap; ws3.cell(i,j).border=bdr
        for j in range(2,4):
            ws3.cell(i,j).font=body
        st=(ws3.cell(i,2).value or "")
        if st=="N/A": ws3.cell(i,2).font=Font(name=F["font"],color="999999",size=10)
        elif st=="Partial": ws3.cell(i,2).font=flag

    # Coverage map: keywords (which kw, tier, placement, why)
    km = ws3.max_row + 2
    ws3.cell(km,1,"COVERAGE MAP — KEYWORDS").font=Font(name=F["font"],bold=True,color="FFFFFF",size=11)
    for col in range(1,7): ws3.cell(km,col).fill=navy
    ws3.append(["Keyword","Tier","Placed In","Why",""," "])
    hr=km+1
    for c in ws3[hr]: c.font=hdr; c.fill=grn; c.border=bdr
    for rowv in d.get("coverage",{}).get("keywords",[]):
        ws3.append(list(rowv)+["",""][:5-len(rowv)]+[""])
    for i in range(hr+1, ws3.max_row+1):
        ws3.cell(i,1).font=bold
        for j in range(1,5):
            ws3.cell(i,j).font = body if j!=1 else bold
            ws3.cell(i,j).alignment=wrap; ws3.cell(i,j).border=bdr

    # Claims to verify
    cv = ws3.max_row + 2
    ws3.cell(cv,1,"CLAIMS TO VERIFY BEFORE PUBLISHING").font=Font(name=F["font"],bold=True,color="FFFFFF",size=11)
    for col in range(1,7): ws3.cell(cv,col).fill=orn
    for claim in d.get("claims_to_verify",[]):
        r=ws3.max_row+1
        ws3.cell(r,1,claim).font=body; ws3.cell(r,1).alignment=wrap
        ws3.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)

    os.makedirs(out_dir, exist_ok=True)
    wb.save(out)
    return out

if __name__ == "__main__":
    draft = sys.argv[1] if len(sys.argv)>1 else "/tmp/listing_draft.json"
    outd  = sys.argv[2] if len(sys.argv)>2 else "/mnt/user-data/outputs"
    print(build(draft, outd))
